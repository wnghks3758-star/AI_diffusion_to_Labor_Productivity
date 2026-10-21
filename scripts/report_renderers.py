from pathlib import Path
import os
import subprocess
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT


def set_run_font(run,size=10.5,bold=False,color=None):
    run.font.name='Noto Sans CJK KR';run.font.size=Pt(size);run.font.bold=bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK KR')
    if color:run.font.color.rgb=RGBColor.from_string(color)

def shade(cell,fill):
    pr=cell._tc.get_or_add_tcPr();v=OxmlElement('w:shd');v.set(qn('w:fill'),fill);pr.append(v)

def word_report(pages,path,title):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
    sec.left_margin=sec.right_margin=Inches(.7);sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.62)
    sec.header_distance=sec.footer_distance=Inches(.28)
    normal=doc.styles['Normal'];normal.font.name='Noto Sans CJK KR';normal.font.size=Pt(11)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.13
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Noto Sans CJK KR')
    header=sec.header.paragraphs[0];set_run_font(header.add_run('Labor Productivity | Continuous DiD · Exploratory analysis'),8,color='667481')
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    set_run_font(footer.add_run('2026.10.21  ·  '),8,color='667481')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    for page_number,page in enumerate(pages):
        if page_number:doc.add_page_break()
        for b in page:
            kind=b['kind']
            if kind in ['title','heading','p','link']:
                p=doc.add_paragraph();p.paragraph_format.keep_with_next=kind in ['title','heading']
                if kind=='title':
                    p.paragraph_format.space_before=Pt(22);p.paragraph_format.space_after=Pt(18)
                    set_run_font(p.add_run(b['text']),25,True,'173F5F')
                elif kind=='heading':
                    p.style=doc.styles['Heading 1'];set_run_font(p.add_run(b['text']),15,True,'173F5F')
                else:
                    text=b['text']+('  '+b['url'] if kind=='link' else '')
                    set_run_font(p.add_run(text),10.5 if len(pages)>10 else 11)
            elif kind=='table':
                table=doc.add_table(rows=1,cols=len(b['headers']));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=False
                widths=np.array(b['widths'],float);widths=widths/widths.sum()*6.85
                for j,width in enumerate(widths):table.columns[j].width=Inches(float(width))
                for row_index,row in enumerate([b['headers']]+b['rows']):
                    cells=table.rows[0].cells if row_index==0 else table.add_row().cells
                    for j,value in enumerate(row):
                        cell=cells[j];cell.width=Inches(float(widths[j]));cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        pr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
                        for side in ['top','bottom','left','right']:
                            e=OxmlElement('w:'+side);e.set(qn('w:w'),'28' if side in ['top','bottom'] else '50');e.set(qn('w:type'),'dxa');margins.append(e)
                        pr.append(margins)
                        p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.04
                        set_run_font(p.add_run(str(value)),b['size'],row_index==0,'FFFFFF' if row_index==0 else '243746')
                        if row_index==0:shade(cell,'245674')
                        elif row_index%2==0:shade(cell,'F1F5F8')
                    rowpr=table.rows[row_index]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');rowpr.append(cant)
                    if row_index==0:
                        repeat=OxmlElement('w:tblHeader');rowpr.append(repeat)
                doc.add_paragraph().paragraph_format.space_after=Pt(0)
            elif kind=='image':
                from PIL import Image
                im=Image.open(b['path']);ratio=im.height/im.width
                width=min(b['width'],b['height']/ratio)
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(b['path']),width=Inches(width))
                p.paragraph_format.space_after=Pt(2)
                if b['caption']:
                    p=doc.add_paragraph();set_run_font(p.add_run(b['caption']),8,color='596A78')
    doc.core_properties.title=title;doc.core_properties.author='Analysis Report'
    doc.core_properties.subject='탐색적 연속 DiD 분석 결과 및 선택 과정'
    doc.save(path)

def esc(text):
    mapping={'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_\allowbreak{}',
             '{':r'\{','}':r'\}','~':r'\textasciitilde{}','^':r'\textasciicircum{}','<':r'\textless{}','>':r'\textgreater{}'}
    return ''.join(mapping.get(c,c) for c in str(text)).replace('\n',r'\\ ')

def pdf_report(pages,stem,title):
    tex=[r'\documentclass[a4paper,11pt]{article}',
         r'\usepackage{fontspec}',r'\setmainfont{Noto Sans CJK KR}',
         r'\XeTeXlinebreaklocale "ko"',r'\XeTeXlinebreakskip=0pt plus 1pt',
         r'\usepackage[top=18mm,bottom=18mm,left=18mm,right=18mm,headheight=12pt]{geometry}',
         r'\usepackage{graphicx,longtable,array,booktabs,xcolor,fancyhdr,hyperref}',
         r'\definecolor{navy}{HTML}{173F5F}',r'\definecolor{muted}{HTML}{596A78}',
         r'\hypersetup{colorlinks=true,urlcolor=navy,pdftitle={'+esc(title)+r'}}',
         r'\pagestyle{fancy}\fancyhf{}',
         r'\fancyhead[L]{\fontsize{8}{10}\selectfont\color{muted} Labor Productivity | Continuous DiD}',
         r'\fancyfoot[R]{\fontsize{8}{10}\selectfont 2026.10.21\quad\thepage}',
         r'\renewcommand{\headrulewidth}{0.3pt}',r'\setlength{\parindent}{0pt}',
         r'\setlength{\parskip}{6pt}',r'\setlength{\tabcolsep}{3pt}',
         r'\setlength{\emergencystretch}{2em}',
         r'\renewcommand{\arraystretch}{1.18}',r'\setlength{\LTpre}{4pt}\setlength{\LTpost}{6pt}',
         r'\begin{document}']
    for i,page in enumerate(pages):
        if i:tex.append(r'\clearpage')
        for b in page:
            kind=b['kind']
            if kind=='title':tex.append(r'{\color{navy}\fontsize{25}{35}\selectfont\bfseries '+esc(b['text'])+r'\par}\vspace{12pt}')
            elif kind=='heading':tex.append(r'{\color{navy}\fontsize{15}{20}\selectfont\bfseries '+esc(b['text'])+r'\par}\vspace{3pt}')
            elif kind=='p':tex.append(r'{\raggedright '+esc(b['text'])+r'\par}'+'\n')
            elif kind=='link':tex.append(esc(b['text'])+r': \url{'+b['url']+r'}\par')
            elif kind=='image':
                tex.append(r'\begin{center}\includegraphics[width='+str(b['width'])+r'in,height='+str(b['height'])+r'in,keepaspectratio]{\detokenize{'+str(b['path'])+r'}}\end{center}')
                if b['caption']:tex.append(r'{\fontsize{8}{10}\selectfont\color{muted} '+esc(b['caption'])+r'\par}')
            elif kind=='table':
                weights=np.array(b['widths'],float);weights/=weights.sum()
                # Account for 6pt inter-column padding in the fixed total text width.
                available=6.85*72-6*len(weights)
                spec=''.join(r'>{\raggedright\arraybackslash}p{'+f'{available*w:.2f}pt'+'}' for w in weights)
                tex += [r'{\fontsize{'+str(b['size'])+'}{'+str(b['size']*1.23)+r'}\selectfont',
                        r'\begin{longtable}{'+spec+'}',r'\toprule',
                        ' & '.join(r'\textbf{'+esc(v)+'}' for v in b['headers'])+r'\\\midrule\endhead']
                for row in b['rows']:tex.append(' & '.join(esc(v) for v in row)+r'\\')
                tex += [r'\bottomrule\end{longtable}}']
    tex.append(r'\end{document}')
    path=WORK/(stem+'.tex');path.write_text('\n'.join(tex))
    log=WORK/(stem+'_compile.txt')
    for _ in range(2):
        run=subprocess.run(['xelatex','-interaction=nonstopmode','-halt-on-error','-output-directory',str(OUT),str(path)],
                           cwd=WORK,env=os.environ.copy(),stdout=subprocess.PIPE,stderr=subprocess.STDOUT,text=True)
        log.write_text(run.stdout)
        if run.returncode:raise RuntimeError('PDF compilation failed: '+str(log))
