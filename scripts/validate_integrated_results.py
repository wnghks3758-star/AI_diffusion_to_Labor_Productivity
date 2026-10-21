"""Independent source, save-gate, adaptive-parent and rendered-table checks."""
from pathlib import Path
import argparse
import json
import sys
import numpy as np
import pandas as pd

ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT))
from dml_acr.io import file_hash,write_json
from dml_acr.config import AnalysisConfig
from dml_acr.random_experiments import PRIMARY, numerical_source_hashes


def validate(root='results/integrated',current='results/adaptive200',report=None):
    root=Path(root);current=Path(current);plan=json.loads((root/'plan.json').read_text())
    local=json.loads((current/'plan.json').read_text());summary=json.loads((root/'summary.json').read_text())
    registry=pd.read_csv(root/'all_specifications_audit.csv');data=pd.read_csv(root/'accepted_estimates.csv')
    cases=[json.loads(line) for line in (root/'retained_cases.jsonl').read_text().splitlines() if line]
    assert len(registry)==len(plan['entries'])==summary['planned']
    assert not registry.run_id.duplicated().any() and not registry.config_id.duplicated().any()
    for item in plan['source_hashes']:
        assert file_hash(item['path'])==item['sha256']
        assert json.loads(Path(item['path']).read_text())['data_sha256']==local['data_sha256']
    for item in local['historical_sources']:assert file_hash(Path(item['root'])/'plan.json')==item['plan_sha256']
    assert file_hash('dml_acr/adaptive_experiments.py')==local['adaptive_source_sha256']
    assert local['numerical_source_hashes']==numerical_source_hashes()
    for name,key in [('anchor_pool.csv','anchor_pool_sha256'),('selected_anchors.csv','selected_anchors_sha256')]:
        assert file_hash(current/name)==local[key]
    pool=pd.read_csv(current/'anchor_pool.csv');selected=pd.read_csv(current/'selected_anchors.csv')
    assert len(pool)==37 and len(selected)==10
    pd.testing.assert_frame_equal(pool.head(10),selected,check_dtype=False)
    prior_ids={e['config_id'] for e in plan['entries'] if e['suite']!=current.name}
    assert len(local['entries'])==200 and len({e['config_id'] for e in local['entries']})==200
    assert len(prior_ids)==864
    parents={r.run_id:r for r in selected.itertuples()}
    for entry in local['entries']:
        assert entry['config_id'] not in prior_ids
        parent=parents[entry['anchor_run_id']]
        original=json.loads((Path(parent.source_folder)/'config.json').read_text())
        assert file_hash(Path(parent.source_folder)/'config.json')==parent.source_config_sha256
        assert file_hash(Path(parent.source_folder)/'estimates.csv')==parent.source_estimates_sha256
        new=entry['config'];normal=original.copy();normal['n_bootstrap']=1999
        changed={}
        for key,value in normal.items():
            if key in ['cache_dir','label']:continue
            if key=='mu_params':
                for param,old in value.items():
                    if old!=new[key][param]:changed['mu_params.'+param]=(old,new[key][param])
            elif value!=new[key]:changed[key]=(value,new[key])
        assert len(changed)==entry['local_coordinate_count'] in [1,2]
        assert set(changed)=={c['coordinate'] for c in entry['changes']}
        for c in entry['changes']:assert changed[c['coordinate']]==(c['before'],c['after'])
        for key in ['gps_features','mu_features']:
            if key in changed:assert len(set(normal[key])^set(new[key]))==1
        assert new['n_bootstrap']==1999
        assert AnalysisConfig.from_dict(new).fingerprint==entry['config_id']
    assert registry.detailed_in_integrated_document.sum()==len(cases)==summary['retained_cases']
    assert set(data.run_id)=={c['run_id'] for c in cases}
    for case in cases:
        source=Path(case['source_folder']);assert file_hash(source/'estimates.csv')==case['source_estimates_sha256']
        table=pd.read_csv(source/'estimates.csv');primary=table.loc[table.effect.isin(PRIMARY)]
        pre=primary.loc[primary.year.isin([2020,2021])]
        assert len(pre)==8 and np.isfinite(pre[['estimate','p_value']].to_numpy()).all() and pre.p_value.ge(.05).all()
        initial=primary.loc[primary.year.eq(2023)&primary.effect.eq('total')].iloc[0]
        assert np.isfinite([initial.estimate,initial.p_value]).all() and initial.estimate<0 and initial.p_value<.05
        combined=data.loc[data.run_id.eq(case['run_id'])].sort_values(['year','effect']).reset_index(drop=True)
        table=table.sort_values(['year','effect']).reset_index(drop=True)
        pd.testing.assert_frame_equal(table[['year','effect']],combined[['year','effect']])
        for column in ['estimate','p_value','se','ci_lower','ci_upper','atet_estimate']:
            np.testing.assert_allclose(table[column],combined[column],rtol=1e-10,atol=1e-14,equal_nan=True)
    selected_rows=data.loc[data.effect.isin(PRIMARY)]
    assert len(selected_rows)==len(cases)*20==summary['verified_primary_rows']
    assert summary['family_tests']==summary['planned']*20
    np.testing.assert_allclose(selected_rows.p_bonferroni_cumulative_budget,
        np.minimum(1,selected_rows.p_value*summary['family_tests']),atol=1e-14,equal_nan=True)
    audit={'passed':True,'planned':summary['planned'],'attempted':summary['attempted'],
        'eligible_cases':len(cases),'verified_primary_rows':len(selected_rows),
        'local_parent_changes_verified':200,'prior_plans_and_source_results_unchanged':True,
        'only_gate_passing_cases_described':True,'cumulative_family_tests':summary['family_tests']}
    if report:
        import fitz
        from docx import Document
        from PIL import Image,ImageDraw
        output=Path(report);stem='Integrated_Search_Report_KO'
        manifest=json.loads((output/'report_validation.json').read_text())
        pdf=fitz.open(output/(stem+'.pdf'));doc=Document(output/(stem+'.docx'))
        tables=[table for table in doc.tables if table.rows[0].cells[0].text=='연도']
        assert len(tables)==len(cases)
        names={'total':'전체','sme':'SME','large':'Large','moderation':'조절효과'}
        def n(value):return f'{float(value):.4f}' if pd.notna(value) else '—'
        def prob(value):return '<.0001' if pd.notna(value) and value<.0001 else n(value)
        count=0
        for case,table in zip(cases,tables):
            sub=data.loc[data.run_id.eq(case['run_id'])&data.effect.isin(PRIMARY)].copy()
            sub['order']=sub.effect.map({e:i for i,e in enumerate(PRIMARY)});sub=sub.sort_values(['year','order'])
            assert len(table.rows)-1==len(sub)==20
            for row,expected in zip(table.rows[1:],sub.itertuples()):
                observed=[cell.text for cell in row.cells]
                target=[str(expected.year),names[expected.effect],n(expected.estimate),
                    f'[{n(expected.ci_lower)}, {n(expected.ci_upper)}]',prob(expected.p_value),prob(expected.p_bonferroni_cumulative_budget)]
                assert observed==target,(case['run_id'],observed,target)
                count+=1
        assert len(pdf)==manifest['report_pages_planned']
        log=(output.parent/'tex'/f'{stem}_compile.txt').read_text()
        assert 'Overfull' not in log and 'Missing character' not in log
        assert all('\ufffd' not in page.get_text() for page in pdf)
        # Neither rejected IDs nor non-gate original cases appear as detailed cases.
        headings=[p.text for p in doc.paragraphs if p.text.startswith('상세 사례')]
        assert len(headings)==2*len(cases)
        audit.update(actual_pdf_pages=len(pdf),docx_numeric_rows_verified=count,layout_overflow_warnings=0)
        manifest.update(audit);write_json(output/'report_validation.json',manifest)
        indices=sorted(set([0,2,4,4+(len(cases)+13)//14,len(pdf)-2,len(pdf)-1]))
        thumbs=[]
        for index in indices:
            pix=pdf[index].get_pixmap(matrix=fitz.Matrix(.65,.65),alpha=False)
            im=Image.frombytes('RGB',(pix.width,pix.height),pix.samples)
            canvas=Image.new('RGB',(im.width,im.height+24),'white');canvas.paste(im,(0,24))
            ImageDraw.Draw(canvas).text((6,4),f'Page {index+1}',fill='black');thumbs.append(canvas)
        canvas=Image.new('RGB',(sum(i.width for i in thumbs),max(i.height for i in thumbs)),'#dddddd');x=0
        for im in thumbs:canvas.paste(im,(x,0));x+=im.width
        canvas.save(output.parent/'contact.png')
    write_json(ROOT/'validation/integrated_audit.json',audit)
    return audit


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--root',default='results/integrated');p.add_argument('--current',default='results/adaptive200');p.add_argument('--report')
    args=p.parse_args();print(json.dumps(validate(args.root,args.current,args.report),ensure_ascii=False))
